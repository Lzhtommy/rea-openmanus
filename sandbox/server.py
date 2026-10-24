"""Sandbox API server — runs INSIDE the sandbox (Docker container or local dev dir).

The agent backend never executes anything itself; every action (shell, file, browser,
port proxy) is an HTTP call to this server. This mirrors Manus' design where the
LLM orchestrator lives outside and a per-task VM exposes an action API.
"""

from __future__ import annotations

import asyncio
import base64
import fnmatch
import os
import re
import signal
import sys
import time
from pathlib import Path

import httpx
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

WORKSPACE = Path(os.environ.get("SANDBOX_WORKSPACE", "/home/ubuntu")).resolve()
WORKSPACE.mkdir(parents=True, exist_ok=True)
SHELL_MAX_BUFFER = 200_000

app = FastAPI(title="manus-sandbox")


def resolve(path: str) -> Path:
    """Resolve a path the agent gave us. Relative paths are relative to the workspace."""
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = WORKSPACE / p
    return p.resolve()


# ---------------------------------------------------------------- shell ---

class ShellSession:
    def __init__(self, sid: str):
        self.id = sid
        self.proc: asyncio.subprocess.Process | None = None
        self.output = ""
        self.command = ""
        self.reader: asyncio.Task | None = None

    @property
    def running(self) -> bool:
        return self.proc is not None and self.proc.returncode is None

    def append(self, text: str) -> None:
        self.output += text
        if len(self.output) > SHELL_MAX_BUFFER:
            self.output = "...[earlier output dropped]...\n" + self.output[-SHELL_MAX_BUFFER:]

    async def start(self, command: str, cwd: Path) -> None:
        self.command = command
        self.output = f"{cwd}$ {command}\n"
        self.proc = await asyncio.create_subprocess_exec(
            "bash", "-lc", command,
            cwd=str(cwd),
            stdin=asyncio.subprocess.PIPE,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.STDOUT,
            start_new_session=True,
            env={**os.environ, "TERM": "dumb", "PAGER": "cat", "GIT_PAGER": "cat"},
        )
        self.reader = asyncio.create_task(self._read())

    async def _read(self) -> None:
        assert self.proc and self.proc.stdout
        while True:
            chunk = await self.proc.stdout.read(4096)
            if not chunk:
                break
            self.append(chunk.decode(errors="replace"))
        rc = await self.proc.wait()
        self.append(f"\n[process exited with code {rc}]\n")

    async def wait(self, seconds: float) -> None:
        if self.reader:
            try:
                await asyncio.wait_for(asyncio.shield(self.reader), timeout=seconds)
            except asyncio.TimeoutError:
                pass

    def kill(self) -> None:
        if self.running:
            try:
                os.killpg(self.proc.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass

    def view(self) -> dict:
        return {
            "id": self.id,
            "command": self.command,
            "running": self.running,
            "returncode": None if self.proc is None else self.proc.returncode,
            "output": self.output,
        }


shells: dict[str, ShellSession] = {}


@app.on_event("shutdown")
def kill_shells() -> None:
    """Shells run in their own sessions; end them with the server. In local mode the sandbox is a
    host process, and a leftover `http.server` would keep its port after the sandbox is gone."""
    for s in shells.values():
        s.kill()


class ShellExec(BaseModel):
    id: str
    command: str
    exec_dir: str | None = None
    timeout: float = 30


@app.post("/shell/exec")
async def shell_exec(req: ShellExec):
    s = shells.get(req.id)
    if s and s.running:
        raise HTTPException(409, f"shell '{req.id}' is busy running `{s.command}`; "
                                 "use shell_wait/shell_view, shell_kill_process, or another shell id")
    s = shells[req.id] = ShellSession(req.id)
    cwd = resolve(req.exec_dir) if req.exec_dir else WORKSPACE
    if not cwd.is_dir():
        raise HTTPException(400, f"exec_dir does not exist: {cwd}")
    await s.start(req.command, cwd)
    await s.wait(req.timeout)
    return s.view()


class ShellRef(BaseModel):
    id: str
    seconds: float = 10
    input: str = ""
    press_enter: bool = True


def _shell(sid: str) -> ShellSession:
    if sid not in shells:
        raise HTTPException(404, f"no shell session '{sid}'")
    return shells[sid]


@app.post("/shell/view")
async def shell_view(req: ShellRef):
    return _shell(req.id).view()


@app.post("/shell/wait")
async def shell_wait(req: ShellRef):
    s = _shell(req.id)
    await s.wait(min(req.seconds, 600))
    return s.view()


@app.post("/shell/write")
async def shell_write(req: ShellRef):
    s = _shell(req.id)
    if not s.running or not s.proc.stdin:
        raise HTTPException(409, "no running process in this shell")
    data = req.input + ("\n" if req.press_enter else "")
    s.append(data)
    s.proc.stdin.write(data.encode())
    await s.proc.stdin.drain()
    await s.wait(2)
    return s.view()


@app.post("/shell/kill")
async def shell_kill(req: ShellRef):
    s = _shell(req.id)
    s.kill()
    await s.wait(2)
    return s.view()


# ----------------------------------------------------------------- file ---

class FileReq(BaseModel):
    file: str
    content: str = ""
    append: bool = False
    leading_newline: bool = False
    trailing_newline: bool = False
    start_line: int | None = None
    end_line: int | None = None
    old_str: str = ""
    new_str: str = ""
    regex: str = ""


@app.post("/file/read")
async def file_read(req: FileReq):
    p = resolve(req.file)
    if not p.is_file():
        raise HTTPException(404, f"file not found: {p}")
    text = p.read_text(errors="replace")
    lines = text.splitlines(keepends=True)
    s = req.start_line or 0
    e = req.end_line if req.end_line is not None else len(lines)
    return {"file": str(p), "content": "".join(lines[s:e]), "total_lines": len(lines)}


@app.post("/file/write")
async def file_write(req: FileReq):
    p = resolve(req.file)
    p.parent.mkdir(parents=True, exist_ok=True)
    content = ("\n" if req.leading_newline else "") + req.content + ("\n" if req.trailing_newline else "")
    with open(p, "a" if req.append else "w") as f:
        f.write(content)
    return {"file": str(p), "bytes": p.stat().st_size}


@app.post("/file/str_replace")
async def file_str_replace(req: FileReq):
    p = resolve(req.file)
    if not p.is_file():
        raise HTTPException(404, f"file not found: {p}")
    text = p.read_text()
    n = text.count(req.old_str)
    if n == 0:
        raise HTTPException(400, "old_str not found in file")
    if n > 1:
        raise HTTPException(400, f"old_str matches {n} times; make it unique")
    p.write_text(text.replace(req.old_str, req.new_str, 1))
    return {"file": str(p), "replaced": 1}


@app.post("/file/find_in_content")
async def file_find_in_content(req: FileReq):
    p = resolve(req.file)
    if not p.is_file():
        raise HTTPException(404, f"file not found: {p}")
    try:
        rx = re.compile(req.regex)
    except re.error as e:
        raise HTTPException(400, f"bad regex: {e}")
    matches = [{"line": i + 1, "text": line.rstrip("\n")}
               for i, line in enumerate(p.read_text(errors="replace").splitlines()) if rx.search(line)]
    return {"file": str(p), "matches": matches[:500], "total": len(matches)}


class FindReq(BaseModel):
    path: str
    glob: str


@app.post("/file/find_by_name")
async def file_find_by_name(req: FindReq):
    root = resolve(req.path)
    if not root.is_dir():
        raise HTTPException(404, f"directory not found: {root}")
    found = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in {"node_modules", ".git", "__pycache__", ".venv"}]
        for name in filenames:
            if fnmatch.fnmatch(name, req.glob):
                found.append(str(Path(dirpath) / name))
                if len(found) >= 1000:
                    return {"files": found, "truncated": True}
    return {"files": found, "truncated": False}


@app.get("/file/download")
async def file_download(path: str):
    p = resolve(path)
    if not p.is_file():
        raise HTTPException(404, f"file not found: {p}")
    return FileResponse(p, filename=p.name)


@app.post("/file/upload")
async def file_upload(file: UploadFile = File(...), dir: str = Form("upload")):
    d = resolve(dir)
    d.mkdir(parents=True, exist_ok=True)
    dest = d / Path(file.filename or "file").name
    dest.write_bytes(await file.read())
    return {"path": str(dest), "bytes": dest.stat().st_size}


# -------------------------------------------------------------- browser ---

# Marks interactive elements in the viewport with data-manus-index and returns a
# compact textual list — the same trick browser-use (which Manus uses) relies on so
# the model can say "click element 12" instead of guessing coordinates.
INDEX_JS = r"""
() => {
  document.querySelectorAll('[data-manus-index]').forEach(e => e.removeAttribute('data-manus-index'));
  const sel = 'a,button,input,textarea,select,summary,[role=button],[role=link],[role=tab],[role=menuitem],[role=checkbox],[onclick],[contenteditable=true]';
  const out = [];
  let i = 0;
  for (const el of document.querySelectorAll(sel)) {
    const r = el.getBoundingClientRect();
    const st = getComputedStyle(el);
    if (r.width < 2 || r.height < 2 || st.visibility === 'hidden' || st.display === 'none') continue;
    if (r.bottom < 0 || r.top > innerHeight || r.right < 0 || r.left > innerWidth) continue;
    el.setAttribute('data-manus-index', String(i));
    let label = (el.innerText || el.value || el.getAttribute('aria-label') || el.getAttribute('placeholder') || el.getAttribute('title') || el.getAttribute('alt') || '').trim().replace(/\s+/g, ' ').slice(0, 80);
    const tag = el.tagName.toLowerCase();
    const typ = el.getAttribute('type');
    out.push(`[${i}]<${tag}${typ ? ' type=' + typ : ''}>${label}</${tag}>`);
    i++;
  }
  return out;
}
"""

TEXT_JS = r"""
() => {
  const sy = scrollY, h = document.documentElement.scrollHeight;
  return {
    text: (document.body ? document.body.innerText : '').replace(/\n{3,}/g, '\n\n'),
    pixels_above: Math.round(sy),
    pixels_below: Math.max(0, Math.round(h - sy - innerHeight)),
  };
}
"""


# With a display (Docker image: Xvfb on :99, exported over VNC) the agent's browser runs headed so the
# user can watch it live and take over; without one (local dev) it runs headless.
HEADED = bool(os.environ.get("DISPLAY"))
# "My Browser": drive the user's own Chrome (started with --remote-debugging-port) instead of launching one,
# so the agent works with the user's logins. Set by the device connector on the user's machine.
CDP_URL = os.environ.get("BROWSER_CDP_URL", "")
SCREEN = {"width": 1280, "height": 900}


class Browser:
    def __init__(self):
        self.pw = None
        self.browser = None
        self.background = None  # headless, for page extraction off-screen
        self.page = None
        self.console: list[str] = []
        self.lock = asyncio.Lock()  # serializes actions on the shared, user-visible page
        self.launch_lock = asyncio.Lock()

    async def _playwright(self):
        try:
            from playwright.async_api import async_playwright
        except ImportError:
            raise HTTPException(501, "playwright not installed in this sandbox")
        if not self.pw:
            self.pw = await async_playwright().start()
        return self.pw

    async def launch(self):
        pw = await self._playwright()
        if CDP_URL and (not self.browser or not self.browser.is_connected()):
            try:
                self.browser = await pw.chromium.connect_over_cdp(CDP_URL)
            except Exception as e:
                raise HTTPException(502, f"cannot reach your Chrome at {CDP_URL}: start it with --remote-debugging-port ({e})")
            return self.browser
        if not self.browser or not self.browser.is_connected():
            args = ["--no-sandbox", "--disable-dev-shm-usage"]
            if HEADED:
                args += ["--window-position=0,0", f"--window-size={SCREEN['width']},{SCREEN['height']}",
                         "--no-first-run", "--no-default-browser-check", "--disable-infobars"]
            self.browser = await pw.chromium.launch(headless=not HEADED, args=args)
        return self.browser

    async def launch_background(self):
        pw = await self._playwright()
        if not self.background or not self.background.is_connected():
            self.background = await pw.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
        return self.background

    async def ensure(self):
        if self.page and not self.page.is_closed():
            return self.page
        async with self.launch_lock:
            await self.launch()
        if CDP_URL:  # a new tab in the user's own browser profile (their cookies and logins)
            ctx = self.browser.contexts[0] if self.browser.contexts else await self.browser.new_context()
            self.page = await ctx.new_page()
            self.page.on("console", lambda m: self.console.append(f"[{m.type}] {m.text}"[:500]))
            return self.page
        # headed: the page fills the window, so what the user sees over VNC is what the agent sees
        ctx = await (self.browser.new_context(no_viewport=True, locale="en-US") if HEADED
                     else self.browser.new_context(viewport={"width": 1280, "height": 800}, locale="en-US"))
        self.page = await ctx.new_page()
        self.page.on("console", lambda m: self.console.append(f"[{m.type}] {m.text}"[:500]))
        return self.page

    async def restart(self):
        try:
            if CDP_URL:  # never close the user's own browser, only the agent's tab
                if self.page and not self.page.is_closed():
                    await self.page.close()
            elif self.browser:
                await self.browser.close()
        except Exception:
            pass
        self.browser = self.page = None
        self.console.clear()

    async def snapshot(self, max_text: int = 12000) -> dict:
        page = self.page
        try:
            await page.wait_for_load_state("domcontentloaded", timeout=5000)
        except Exception:
            pass
        elements = await page.evaluate(INDEX_JS)
        info = await page.evaluate(TEXT_JS)
        shot = await page.screenshot(type="jpeg", quality=60)
        text = info["text"]
        if len(text) > max_text:
            text = text[:max_text] + f"\n...[page text truncated, {len(info['text'])} chars total; scroll or use browser_console_exec to read more]"
        return {
            "url": page.url,
            "title": await page.title(),
            "elements": elements,
            "text": text,
            "pixels_above": info["pixels_above"],
            "pixels_below": info["pixels_below"],
            "screenshot": base64.b64encode(shot).decode(),
        }


browser = Browser()


class BrowserReq(BaseModel):
    url: str = ""
    index: int | None = None
    coordinate_x: float | None = None
    coordinate_y: float | None = None
    text: str = ""
    press_enter: bool = False
    key: str = ""
    option: int = 0
    javascript: str = ""
    to_end: bool = False


async def _locate(page, req: BrowserReq):
    if req.index is not None:
        loc = page.locator(f'[data-manus-index="{req.index}"]')
        if await loc.count() == 0:
            raise HTTPException(400, f"element index {req.index} not found; call browser_view to refresh indexes")
        return loc.first
    return None


@app.post("/browser/{action}")
async def browser_action(action: str, req: BrowserReq):
    async with browser.lock:
        if action == "restart":
            await browser.restart()
            page = await browser.ensure()
            if req.url:
                await page.goto(req.url, timeout=45000, wait_until="domcontentloaded")
            return await browser.snapshot()
        page = await browser.ensure()
        try:
            if action == "navigate":
                await page.goto(req.url, timeout=45000, wait_until="domcontentloaded")
            elif action == "view":
                pass
            elif action == "click":
                loc = await _locate(page, req)
                if loc:
                    await loc.click(timeout=8000)
                elif req.coordinate_x is not None:
                    await page.mouse.click(req.coordinate_x, req.coordinate_y or 0)
                else:
                    raise HTTPException(400, "provide index or coordinates")
                await page.wait_for_timeout(800)
            elif action == "input":
                loc = await _locate(page, req)
                if loc:
                    await loc.fill(req.text, timeout=8000)
                    if req.press_enter:
                        await loc.press("Enter")
                elif req.coordinate_x is not None:
                    await page.mouse.click(req.coordinate_x, req.coordinate_y or 0)
                    await page.keyboard.type(req.text)
                    if req.press_enter:
                        await page.keyboard.press("Enter")
                else:
                    raise HTTPException(400, "provide index or coordinates")
                await page.wait_for_timeout(800)
            elif action == "move_mouse":
                await page.mouse.move(req.coordinate_x or 0, req.coordinate_y or 0)
            elif action == "press_key":
                await page.keyboard.press(req.key)
                await page.wait_for_timeout(500)
            elif action == "select_option":
                loc = await _locate(page, req)
                if not loc:
                    raise HTTPException(400, "index required")
                await loc.select_option(index=req.option)
            elif action in ("scroll_down", "scroll_up"):
                sign = 1 if action == "scroll_down" else -1
                if req.to_end:
                    await page.evaluate(f"window.scrollTo(0, {sign} > 0 ? document.documentElement.scrollHeight : 0)")
                else:
                    await page.evaluate(f"window.scrollBy(0, {sign} * window.innerHeight * 0.85)")
                await page.wait_for_timeout(400)
            elif action == "console_exec":
                result = await page.evaluate(req.javascript)
                snap = await browser.snapshot()
                snap["console_result"] = repr(result)[:20000]
                return snap
            elif action == "console_view":
                return {"url": page.url, "console": browser.console[-200:]}
            else:
                raise HTTPException(404, f"unknown browser action {action}")
        except HTTPException:
            raise
        except Exception as e:  # surface playwright errors to the model verbatim
            raise HTTPException(400, f"{type(e).__name__}: {e}")
        return await browser.snapshot()


# ------------------------------------------------------------ web extract ---
# Read-only page fetch in a throwaway context (Manus' `webpage_extract`): does not touch
# the shared browser page, so several extracts can run concurrently.

class ExtractReq(BaseModel):
    url: str
    max_chars: int = 30000


@app.post("/web/extract")
async def web_extract(req: ExtractReq):
    async with browser.launch_lock:
        b = await browser.launch_background()
    ctx = await b.new_context(viewport={"width": 1280, "height": 800}, locale="en-US")
    try:
        page = await ctx.new_page()
        try:
            await page.goto(req.url, timeout=45000, wait_until="domcontentloaded")
            await page.wait_for_timeout(500)
            info = await page.evaluate(TEXT_JS)
            title = await page.title()
        except Exception as e:
            raise HTTPException(400, f"{type(e).__name__}: {e}")
        text = info["text"].strip()
        return {"url": page.url, "title": title, "text": text[:req.max_chars],
                "truncated": len(text) > req.max_chars, "total_chars": len(text)}
    finally:
        await ctx.close()


class ShotReq(BaseModel):
    url: str  # http(s) URL or a sandbox path of an HTML file
    width: int = 1280
    height: int = 800
    wait_ms: int = 1200


@app.post("/web/screenshot")
async def web_screenshot(req: ShotReq):
    """A JPEG of a page in a separate headless tab (e.g. a cover image for a published site)."""
    url = req.url if req.url.startswith(("http://", "https://", "file://")) else resolve(req.url).as_uri()
    async with browser.launch_lock:
        b = await browser.launch_background()
    ctx = await b.new_context(viewport={"width": req.width, "height": req.height}, locale="zh-CN")
    try:
        page = await ctx.new_page()
        try:
            await page.goto(url, timeout=45000, wait_until="load")
            await page.wait_for_timeout(req.wait_ms)
            jpg = await page.screenshot(type="jpeg", quality=75)
        except Exception as e:
            raise HTTPException(400, f"{type(e).__name__}: {e}")
        return {"image": base64.b64encode(jpg).decode(), "title": await page.title()}
    finally:
        await ctx.close()


# ---------------------------------------------------------------- recorder ---
# Skill recorder: while the user demonstrates a workflow in the browser, a page script reports what they
# do (clicks, typing, submits, navigation) with a human-readable description of each target element.
# Passwords are never captured.

RECORDER_JS = r"""
(() => {
  if (window.__omRecorder) return; window.__omRecorder = true;
  const label = (el) => {
    if (!el || !el.tagName) return '';
    const t = (s) => (s || '').replace(/\s+/g, ' ').trim().slice(0, 80);
    const parts = [el.tagName.toLowerCase()];
    const name = el.getAttribute('aria-label') || el.getAttribute('title') || el.getAttribute('placeholder') || el.getAttribute('name') || el.getAttribute('alt');
    const text = /^(INPUT|TEXTAREA)$/.test(el.tagName) ? '' : t(el.innerText || '');  // never a field's value
    if (el.id) parts.push('#' + el.id);
    if (name) parts.push(`"${t(name)}"`);
    if (text && text !== name) parts.push(`text="${text.slice(0, 60)}"`);
    if (el.type && el.tagName === 'INPUT') parts.push(`type=${el.type}`);
    return parts.join(' ');
  };
  const target = (e) => e.target.closest('a,button,input,textarea,select,[role=button],[role=link],[role=tab],[role=menuitem],label,summary') || e.target;
  const send = (ev) => { try { window.__omRecord(JSON.stringify({ ...ev, url: location.href, title: document.title, t: Date.now() })); } catch (_) {} };
  document.addEventListener('click', (e) => { const el = target(e); send({ type: 'click', target: label(el), href: el.href || '' }); }, true);
  document.addEventListener('change', (e) => {
    const el = e.target, secret = el.type === 'password' || /pass|pwd|token|secret|otp|code/i.test(el.name || el.id || '');
    const value = el.tagName === 'SELECT' ? (el.selectedOptions[0]?.text || el.value) : el.type === 'checkbox' ? String(el.checked) : el.value;
    send({ type: 'input', target: label(el), value: secret ? '••••' : String(value).slice(0, 300), secret });
  }, true);
  document.addEventListener('keydown', (e) => { if (e.key === 'Enter' && e.target.tagName !== 'TEXTAREA') send({ type: 'key', key: 'Enter', target: label(e.target) }); }, true);
  document.addEventListener('submit', (e) => send({ type: 'submit', target: label(e.target) }), true);
})();
"""


class Recorder:
    def __init__(self):
        self.events: list[dict] = []
        self.active = False
        self.bound: set[int] = set()

    async def _bind(self, page) -> None:
        if id(page.context) not in self.bound:
            self.bound.add(id(page.context))
            await page.context.expose_binding("__omRecord", lambda src, raw: self._on(raw))
            await page.context.add_init_script(RECORDER_JS)
            page.context.on("page", lambda p: p.on("framenavigated", lambda f: self._nav(p, f)))
            page.on("framenavigated", lambda f: self._nav(page, f))
        await page.evaluate(RECORDER_JS)

    def _on(self, raw: str) -> None:
        if self.active:
            try:
                self.events.append(__import__("json").loads(raw))
            except ValueError:
                pass

    def _nav(self, page, frame) -> None:
        if self.active and frame == page.main_frame and frame.url not in ("about:blank",):
            if not self.events or self.events[-1].get("url") != frame.url or self.events[-1]["type"] != "navigate":
                self.events.append({"type": "navigate", "url": frame.url, "t": int(time.time() * 1000)})


recorder = Recorder()


class RecordReq(BaseModel):
    url: str = ""


@app.post("/browser/record/start")
async def record_start(req: RecordReq):
    async with browser.lock:
        page = await browser.ensure()
        recorder.events, recorder.active = [], True
        await recorder._bind(page)
        if req.url:
            try:
                await page.goto(req.url, timeout=45000, wait_until="domcontentloaded")
            except Exception as e:  # keep recording: the user can navigate by hand
                return {"ok": True, "url": page.url, "warning": f"could not open {req.url}: {e}"[:300]}
    return {"ok": True, "url": page.url}


@app.post("/browser/record/stop")
async def record_stop():
    recorder.active = False
    return {"events": recorder.events[-500:], "count": len(recorder.events)}


@app.get("/browser/record/events")
async def record_events():
    return {"events": recorder.events[-50:], "count": len(recorder.events), "active": recorder.active}


# ---------------------------------------------------------------- slides ---
# An HTML deck is a page of <section class="slide"> elements designed at a fixed size (default
# 1280x720). Export renders each slide on its own: PNG images (previews and the PPTX), a vector PDF
# (print CSS, one page per slide) and a PPTX with one full-bleed image per slide plus speaker notes.

class SlidesReq(BaseModel):
    file: str
    formats: list[str] = ["pdf", "pptx"]
    width: int = 1280
    height: int = 720


SLIDE_SEL = "section.slide, .slide"
ISOLATE_JS = """([sel, i, w, h]) => {
  const all = [...document.querySelectorAll(sel)].filter(s => !s.parentElement.closest(sel));
  document.documentElement.classList.add('exporting');
  all.forEach((s, k) => { s.style.display = k === i ? '' : 'none'; s.classList.toggle('active', k === i); });
  const s = all[i];
  if (getComputedStyle(s).display === 'none') s.style.display = 'block';  // decks that show slides by other means
  Object.assign(s.style, {position: 'fixed', left: '0', top: '0', width: w + 'px', height: h + 'px', margin: '0',
                          transform: 'none', zIndex: '2147483647', overflow: 'hidden'});
  const n = s.querySelector('aside.notes, .notes');
  return {title: (s.querySelector('h1,h2,h3')?.innerText || '').trim().slice(0, 120), notes: n ? n.innerText.trim() : ''};
}"""
PRINT_CSS = """@page { size: %dpx %dpx; margin: 0 }
html, body { margin: 0 !important; padding: 0 !important; background: none !important; width: %dpx !important; }
.deck, .slides { transform: none !important; display: block !important; width: auto !important; height: auto !important; }
%s { position: relative !important; left: 0 !important; top: 0 !important;
  width: %dpx !important; height: %dpx !important; margin: 0 !important; transform: none !important;
  break-after: page; page-break-after: always; overflow: hidden !important; opacity: 1 !important; visibility: visible !important; }
aside.notes, .notes, .deck-ui { display: none !important; }"""


@app.post("/slides/export")
async def slides_export(req: SlidesReq):
    src = resolve(req.file)
    if not src.is_file():
        raise HTTPException(404, f"file not found: {src}")
    w, h = req.width, req.height
    out_dir = src.parent / f"{src.stem}_slides"
    out_dir.mkdir(exist_ok=True)
    async with browser.launch_lock:
        b = await browser.launch_background()
    ctx = await b.new_context(viewport={"width": w, "height": h}, device_scale_factor=1.5, locale="zh-CN")
    try:
        page = await ctx.new_page()
        await page.goto(src.as_uri(), wait_until="load", timeout=60000)
        await page.wait_for_timeout(800)  # web fonts, charts
        count = await page.evaluate(f"() => [...document.querySelectorAll('{SLIDE_SEL}')].filter(s => !s.parentElement.closest('{SLIDE_SEL}')).length")
        if not count:
            raise HTTPException(400, "no slides found: wrap each slide in <section class=\"slide\">")
        slides, thumbs = [], []
        for i in range(count):
            meta = await page.evaluate(ISOLATE_JS, [SLIDE_SEL, i, w, h])
            await page.wait_for_timeout(120)
            png = await page.screenshot(clip={"x": 0, "y": 0, "width": w, "height": h})
            path = out_dir / f"slide-{i + 1:02d}.png"
            path.write_bytes(png)
            jpg = await page.screenshot(clip={"x": 0, "y": 0, "width": w, "height": h}, type="jpeg", quality=70, scale="css")
            thumbs.append(base64.b64encode(jpg).decode())
            slides.append({**meta, "image": str(path)})
        out = {"count": count, "slides": slides, "thumbnails": thumbs}
        if "pdf" in req.formats:
            pdf_page = await ctx.new_page()
            await pdf_page.goto(src.as_uri(), wait_until="load", timeout=60000)
            await pdf_page.add_style_tag(content=PRINT_CSS % (w, h, w, SLIDE_SEL, w, h))
            await pdf_page.evaluate(f"""() => {{ document.documentElement.classList.add('exporting');
              document.querySelectorAll('{SLIDE_SEL}').forEach(s => {{ s.classList.add('active');
                if (getComputedStyle(s).display === 'none') s.style.display = 'block'; }}); }}""")
            await pdf_page.wait_for_timeout(500)
            pdf = src.with_suffix(".pdf")
            await pdf_page.pdf(path=str(pdf), width=f"{w}px", height=f"{h}px", print_background=True, prefer_css_page_size=True)
            out["pdf"] = str(pdf)
        if "pptx" in req.formats:
            out["pptx"] = await asyncio.to_thread(_build_pptx, slides, src.with_suffix(".pptx"), w, h)
        return out
    finally:
        await ctx.close()


def _build_pptx(slides: list[dict], dest: Path, w: int, h: int) -> str:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(w * 9525), Emu(h * 9525)  # 1 px = 9525 EMU at 96 dpi
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(s["image"], 0, 0, width=prs.slide_width, height=prs.slide_height)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(str(dest))
    return str(dest)


# ---------------------------------------------------------------- proxy ---
# Lets the backend expose a service the agent started on an internal port
# (deploy_expose_port) without publishing every container port.

proxy_client = httpx.AsyncClient(timeout=60)


@app.api_route("/proxy/{port}/{path:path}", methods=["GET", "POST", "PUT", "DELETE", "PATCH", "HEAD", "OPTIONS"])
async def proxy(port: int, path: str, request: Request):
    url = f"http://127.0.0.1:{port}/{path}"
    if request.url.query:
        url += "?" + request.url.query
    headers = {k: v for k, v in request.headers.items() if k.lower() not in {"host", "content-length"}}
    try:
        r = await proxy_client.request(request.method, url, headers=headers, content=await request.body())
    except httpx.HTTPError as e:
        raise HTTPException(502, f"nothing reachable on port {port}: {e}")
    drop = {"content-encoding", "transfer-encoding", "connection", "content-length"}
    return Response(r.content, status_code=r.status_code,
                    headers={k: v for k, v in r.headers.items() if k.lower() not in drop})


# ------------------------------------------------------------------ vnc ---
# WebSocket <-> TCP bridge to x11vnc (localhost:5900), so the live screen goes through the
# same single port as the action API. The orchestrator relays it to the web client (noVNC).

VNC_PORT = 5900


async def vnc_available() -> bool:
    try:
        _, w = await asyncio.wait_for(asyncio.open_connection("127.0.0.1", VNC_PORT), 1)
        w.close()
        return True
    except (OSError, asyncio.TimeoutError):
        return False


@app.websocket("/vnc")
async def vnc(ws: WebSocket):
    proto = "binary" if "binary" in ws.scope.get("subprotocols", []) else None
    try:
        reader, writer = await asyncio.open_connection("127.0.0.1", VNC_PORT)
    except OSError:
        await ws.close(code=1011)
        return
    await ws.accept(subprotocol=proto)

    async def up():
        try:
            while True:
                msg = await ws.receive()
                if msg["type"] == "websocket.disconnect":
                    break
                data = msg.get("bytes") or (msg.get("text") or "").encode()
                writer.write(data)
                await writer.drain()
        except (WebSocketDisconnect, ConnectionError):
            pass

    async def down():
        try:
            while chunk := await reader.read(65536):
                await ws.send_bytes(chunk)
        except (WebSocketDisconnect, ConnectionError, RuntimeError):
            pass

    tasks = [asyncio.create_task(up()), asyncio.create_task(down())]
    await asyncio.wait(tasks, return_when=asyncio.FIRST_COMPLETED)
    for t in tasks:
        t.cancel()
    writer.close()
    try:
        await ws.close()
    except RuntimeError:
        pass


@app.get("/health")
async def health():
    return {"ok": True, "workspace": str(WORKSPACE), "time": time.time(), "vnc": await vnc_available(),
            "screen": SCREEN if HEADED else None, "browser": "cdp" if CDP_URL else "chromium",
            "platform": sys.platform, "host": os.uname().nodename if hasattr(os, "uname") else ""}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host=os.environ.get("SANDBOX_HOST", "0.0.0.0"), port=int(os.environ.get("SANDBOX_PORT", "8080")))
